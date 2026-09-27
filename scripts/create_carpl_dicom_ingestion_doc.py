#!/usr/bin/env python3
"""Generate the editable CARPL DICOM-ingestion technical note."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "CARPL_RAPIER_BOX_MRI_Modality_Filtering_and_Ingestion_Guide.docx"

NAVY = "17365D"
BLUE = "2B6CB0"
TEAL = "138A8A"
GREEN = "2F855A"
AMBER = "B7791F"
RED = "B83232"
LIGHT_BLUE = "EAF2F8"
LIGHT_TEAL = "E8F6F3"
LIGHT_AMBER = "FFF8E1"
LIGHT_RED = "FDECEC"
LIGHT_GREY = "F3F5F7"
MID_GREY = "D7DEE5"
WHITE = "FFFFFF"
BLACK = "1A202C"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 28, NAVY),
        ("Subtitle", 13, TEAL),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 10.5, TEAL),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Aptos Mono"
    code.font.size = Pt(7.2)
    code.font.color.rgb = RGBColor.from_string(BLACK)
    code.paragraph_format.left_indent = Cm(0.35)
    code.paragraph_format.right_indent = Cm(0.2)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "SGH Prostate MRI AI  |  CARPL / RAPIER BOX Integration Note"
    header.style = styles["Caption"]
    header.runs[0].font.color.rgb = RGBColor.from_string(NAVY)
    header.runs[0].font.size = Pt(8)
    add_page_number(section.footer.paragraphs[0])


def add_text(doc: Document, text: str, *, bold=False, italic=False, color=None, style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def add_bullets(doc: Document, items, level=0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        if isinstance(item, tuple):
            lead, rest = item
            p.add_run(lead).bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_numbered(doc: Document, items) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        if isinstance(item, tuple):
            lead, rest = item
            p.add_run(lead).bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_callout(doc: Document, title: str, body: str, *, kind="info") -> None:
    palette = {
        "info": (BLUE, LIGHT_BLUE),
        "safe": (GREEN, LIGHT_TEAL),
        "warn": (AMBER, LIGHT_AMBER),
        "risk": (RED, LIGHT_RED),
    }
    accent, fill = palette[kind]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_cell_width(table.cell(0, 0), 0.12)
    set_cell_width(table.cell(0, 1), 6.75)
    set_cell_shading(table.cell(0, 0), accent)
    set_cell_shading(table.cell(0, 1), fill)
    set_cell_margins(table.cell(0, 0), 0, 0, 0, 0)
    set_cell_margins(table.cell(0, 1), 110, 160, 110, 160)
    p = table.cell(0, 1).paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.add_run(title + "\n").bold = True
    p.runs[0].font.color.rgb = RGBColor.from_string(accent)
    p.add_run(body)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc: Document, headers, rows, widths=None, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        if widths:
            set_cell_width(cell, widths[i])
        p = cell.paragraphs[0]
        r = p.add_run(str(header))
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        r.font.size = Pt(font_size)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, value in enumerate(row):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if widths:
                set_cell_width(cell, widths[i])
            if ridx % 2:
                set_cell_shading(cell, LIGHT_GREY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc: Document, code: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_shading(table.cell(0, 0), LIGHT_GREY)
    set_cell_margins(table.cell(0, 0), 110, 130, 110, 130)
    p = table.cell(0, 0).paragraphs[0]
    p.style = doc.styles["Code Block"]
    p.add_run(code)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_flow_table(doc: Document) -> None:
    headers = ["1. CARPL / PACS", "2. RAPIER input", "3. In-box selection", "4. AI pipeline", "5. Return / review"]
    rows = [[
        "Identify one MR study; classify complete series; reject ambiguity",
        "Read-only case folder containing exactly one T2, ADC and high-b series",
        "Recursive DICOM load → three regex selectors → max-b instance filter",
        "T2 organ model → 3-channel lesion ensemble → PI-RADS classifier",
        "NIfTI, RTSTRUCT, DICOM SR, report and copies of selected DICOMs",
    ]]
    add_table(doc, headers, rows, widths=[1.35] * 5, font_size=7.5)


def build_document() -> Path:
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "CARPL / RAPIER BOX MRI Modality Filtering and Ingestion Guide"
    props.subject = "Code-backed T2, ADC and high-b DICOM selection and inference pipeline"
    props.author = "SGH prostate MRI AI project"
    props.keywords = "CARPL, RAPIER BOX, DICOM, MRI, T2, ADC, high-b, MONAI Deploy"

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(46)
    p.add_run("TECHNICAL INTEGRATION NOTE").bold = True
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    title = doc.add_paragraph(style="Title")
    title.add_run("MRI modality filtering and ingestion\nfor the CARPL → RAPIER BOX workflow")
    sub = doc.add_paragraph(style="Subtitle")
    sub.add_run("A code-backed guide to selecting T2, ADC and high-b DWI for the SGH prostate MRI lesion application")
    doc.add_paragraph()
    add_callout(
        doc,
        "Purpose",
        "Give the CARPL engineering team an implementation-level contract for sending the minimum correct DICOM content, understanding what the RAPIER BOX app validates locally, and customising series rules without silently changing the model input distribution.",
        kind="info",
    )
    add_table(
        doc,
        ["Document status", "Repository basis", "Prepared", "Audience"],
        [["Editable working document", "Current checked-out source; app.py is authoritative", date.today().isoformat(), "CARPL integration, SGH radiology IT, AI operations"]],
        widths=[1.55, 2.25, 1.05, 1.9],
        font_size=8,
    )
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Scope note. ").bold = True
    p.add_run("This describes inference and DICOM routing. Training code was reviewed only where needed to confirm the deployed input contract; CARPL is not expected to run the training workflow.")
    p = doc.add_paragraph()
    p.add_run("Clinical status. ").bold = True
    p.add_run("The source application identifies itself as non-commercial research software and states that it must not be used for treatment or diagnosis. AI results require radiologist review.")

    page_break(doc)
    doc.add_heading("Contents", level=1)
    contents = [
        "1. Executive integration decision",
        "2. What the application actually receives",
        "3. Exact T2, ADC and HIGHB selection logic",
        "4. Highest-b-value instance filtering",
        "5. DICOM geometry and volume construction",
        "6. End-to-end inference pipeline",
        "7. Software, libraries, models and hardware",
        "8. Storage-efficient CARPL hand-off design",
        "9. Customisation method and change control",
        "10. Validation, monitoring and failure handling",
        "11. CARPL implementation checklist",
        "Appendix A. Exact deployed regular expressions",
        "Appendix B. Required tags and preservation rules",
        "Appendix C. Source-code traceability map",
    ]
    add_numbered(doc, contents)
    add_callout(doc, "Reading priority", "Sections 1, 3, 4, 8 and 11 contain the interface contract CARPL should implement. Sections 5–7 explain why the contract matters to model behavior.", kind="safe")

    page_break(doc)
    doc.add_heading("1. Executive integration decision", level=1)
    add_callout(
        doc,
        "Recommended contract",
        "CARPL should stage one and only one complete T2 series, one complete ADC series, and one high-b DWI series from the same MR study. For a multi-b DWI Series Instance UID, stage only instances at the highest numeric b-value after validating slice completeness. Preserve original identifiers and geometry. The RAPIER BOX should still keep its local selectors as a defensive check.",
        kind="safe",
    )
    add_flow_table(doc)
    doc.add_heading("Why this is the safest design", level=2)
    add_bullets(doc, [
        ("Correctness: ", "the lesion network was built for three registered channels in the order T2, ADC, HIGHB; a wrong route creates a plausible-looking but invalid model input."),
        ("Determinism: ", "the current MONAI selector returns the first matching series. Sending exactly one candidate per role removes filesystem/order dependence."),
        ("Storage: ", "local staging does not need the complete multiparametric study once CARPL has selected and validated the three required series."),
        ("Traceability: ", "the T2 Series Instance UID is used as the reference for RTSTRUCT and DICOM SR evidence."),
        ("Defence in depth: ", "matching rules remain in the app, so an incorrectly routed or renamed series can still be rejected before inference."),
    ])
    doc.add_heading("What must not be inferred", level=2)
    add_bullets(doc, [
        "A folder name such as “T2” is not used by the model to classify a series; the deployed selectors read DICOM metadata.",
        "HIGHB selection is not initially driven by Diffusion b-value tags. SeriesDescription first identifies the diffusion candidate; b-value tags then prune instances inside that series.",
        "The code chooses the maximum available b-value; it does not enforce a clinical minimum such as b=1400 or b=1800.",
        "The app does not currently score image quality, prostate coverage, motion, repeat-series preference, or newest acquisition time.",
    ])

    doc.add_heading("2. What the application actually receives", level=1)
    add_text(doc, "At runtime the Docker wrapper mounts one case directory read-only at /input, an output directory at /output, and the model tree read-only at /models. The MONAI DICOM loader recursively enumerates every file beneath /input, ignores files that pydicom declares non-DICOM, groups valid instances by Study Instance UID and Series Instance UID, and then emits in-memory studies and series.")
    add_table(doc, ["Boundary", "Current behavior", "CARPL consequence"], [
        ["Input unit", "A filesystem directory; nesting is allowed because loading is recursive.", "Create a separate staging folder/job per patient study or accession."],
        ["Study scope", "More than one study can be loaded; no app-level assertion requires the three selected roles to share one Study UID.", "Enforce exactly one Study Instance UID before transfer."],
        ["Series scope", "Each of the three selectors independently returns the first matching series per study; the converter then consumes the first selected study/series.", "Reject zero or multiple candidates upstream; do not rely on ordering."],
        ["Files", "All files are visited. Valid DICOM is parsed and non-DICOM is ignored.", "Sending unrelated series costs I/O and storage and expands the ambiguity surface."],
        ["Outputs", "The app writes NIfTI intermediates/results and copies the selected source DICOM series under output/dicom/.", "Plan retention to avoid keeping both a redundant staging copy and exported copy indefinitely."],
    ], widths=[1.15, 2.8, 2.8])
    add_callout(doc, "Important current limitation", "There is no cross-modality Study UID or Frame of Reference UID equality gate in app.py. ADC and HIGHB are resampled to T2, but resampling cannot make unrelated acquisitions clinically registered. CARPL should treat same-study, same-patient and geometry compatibility as hard preconditions.", kind="risk")

    page_break(doc)
    doc.add_heading("3. Exact T2, ADC and HIGHB selection logic", level=1)
    add_text(doc, "All three selectors require DICOM Modality (0008,0060) to match MR. SeriesDescription (0008,103E) is matched using case-insensitive regular-expression search. ImageType is intentionally not required because the SGH de-identified data frequently lacks it.")
    add_table(doc, ["Role", "Positive conditions", "Exclusions", "Examples from the annotated SGH set"], [
        ["T2", "Contains t2 AND an axial marker: tra, standalone ax, axial, or an ax_ / _ax_ pattern.", "Any marker for oblique, coronal, sagittal, pelvis/whole pelvis, TIRM, STIR, SPACE, localizer, survey, T1, +C, ADC, DWI/diff/TRACEW, or selected FS tokens.", "t2_tse_tra_BLADE PROSTATE; Ax T2 PROP; AX_t2_tse"],
        ["ADC", "Contains adc.", "localizer or survey.", "…prostate_ADC; ADC (10^-6 mm²/s); ddep2d_diff_ADC"],
        ["HIGHB", "Contains tracew, dwi, diff, calc, compute/computed/computer; or focus/muse followed later by b plus a digit.", "adc, localizer or survey.", "…prostate_TRACEW; Ax FOCUS b1800; calculated b1800"],
    ], widths=[0.65, 2.15, 2.2, 1.75], font_size=7.8)
    doc.add_heading("Selection precedence and overlaps", level=2)
    add_bullets(doc, [
        "An ADC description may also contain DWI, DIFF or TRACEW; the explicit HIGHB exclusion for adc ensures that these ADC maps route to ADC and not HIGHB.",
        "The T2 expression excludes ADC and diffusion markers, so a mixed or misleading description is not treated as anatomical T2.",
        "The selector configuration uses all_matched=False (the default). If two T2 series both match, only the first encountered match is returned. The same applies independently to ADC and HIGHB.",
        "SeriesDescription is a mandatory practical input: an absent/blank description cannot satisfy these expressions and therefore produces no match.",
    ])
    doc.add_heading("Evidence for these expressions", level=2)
    add_text(doc, "The repository includes a validation harness that extracts the expressions directly from app.py and applies MONAI-equivalent case-insensitive matching to a manually annotated SGH CSV. Re-running it for this document evaluated 43,470 rows: 3,562 T2, 3,543 ADC, 4,502 HIGHB and 31,863 ignore rows. It reported zero false positives and zero false negatives for all three classes.")
    add_callout(doc, "How to interpret 100%", "This is strong retrospective evidence on the annotated SGH naming distribution, not proof for future scanner protocols, sites, vendor upgrades, edited SeriesDescriptions, or DICOMweb transformations. Every rule change needs a new representative labelled validation set.", kind="warn")

    doc.add_heading("4. Highest-b-value instance filtering", level=1)
    add_text(doc, "After a HIGHB series is selected by description, HighBValueFilterOperator inspects every SOP instance. This is necessary because SGH acquisitions may place multiple diffusion weightings in one Series Instance UID; stacking all of them as one z-volume would interleave b-values and produce a physically meaningless 3-D input.")
    add_numbered(doc, [
        ("Read b-value. ", "Use standard Diffusion b-value (0018,9087). If absent/unusable, try Siemens private (0019,100C)."),
        ("Normalise. ", "Convert to float and round to the nearest integer, so 1799.99 is compared as 1800."),
        ("Find maximum. ", "Collect known values and choose max(known)."),
        ("Prune when genuinely multi-b. ", "If two or more distinct known values exist, retain only instances whose rounded value equals the maximum; lower-b and untagged instances are removed."),
        ("Pass through special cases. ", "If no tag is found anywhere, or only one distinct known value is present, the implementation leaves the entire series unchanged."),
    ])
    add_table(doc, ["Observed tag state", "Current in-box action", "Recommended CARPL action"], [
        ["Multiple known b-values", "Keep only instances at the numeric maximum.", "Do the same, then verify a complete spatial stack at that b-value."],
        ["Exactly one known b-value", "No pruning; even untagged instances remain.", "Confirm every image belongs to that b-value and flag mixed tagged/untagged content."],
        ["No recognised b-value tags", "Pass through the entire matched diffusion series.", "Do not guess. Quarantine/manual review or add a scanner-specific, validated extractor."],
        ["Private tag other than Siemens 0019,100C", "Not recognised by current code.", "Map only after vendor-specific validation; record the source tag and value."],
    ], widths=[1.45, 2.5, 2.6])
    add_callout(doc, "Storage example from the repository test data", "One SGH test case contains 54 DICOM files across 10 series. The matching candidates are 6 T2 instances, 1 ADC instance and a 21-instance multi-b diffusion series. The latter contains b=50/500/1000/1800 and is reduced to 7 b=1800 instances. An upstream minimum payload would therefore be 14 images rather than all 54—but the 1-slice ADC and thin stacks should be flagged for image-quality review, not treated as a production-quality example.", kind="info")

    page_break(doc)
    doc.add_heading("5. DICOM geometry and volume construction", level=1)
    add_text(doc, "MONAI Deploy converts each selected series into an in-memory 3-D image. The app applies a local patch to the upstream volume preparation because SGH de-identified DICOM exposed incorrect slice-removal and sorting behavior.")
    add_table(doc, ["Step", "Code behavior", "Operational meaning"], [
        ["Spatial eligibility", "For each SOP, read ImageOrientationPatient (0020,0037) and ImagePositionPatient (0020,0032).", "Both are needed to compute a slice normal and physical distance."],
        ["Bad slice handling", "Remove instances missing either orientation or position. If all are removed, raise ValueError.", "Do not send incomplete spatial stacks expecting InstanceNumber to repair them."],
        ["Ordering", "Sort retained slices by physical distance; a fallback can sort by InstanceNumber if distance is unavailable.", "Physical geometry, not filename order, defines the volume."],
        ["Through-plane spacing", "Compute distance between the first two retained slice positions. A one-instance series receives depth spacing 1.0.", "A single-slice series can technically convert but is not a valid 3-D prostate input."],
        ["Affine", "Create DICOM/NIfTI affine metadata. Later code falls back from nifti_affine_transform to dicom_affine_transform, then identity with a warning.", "Identity fallback keeps execution alive but makes spatial reliability poor."],
        ["Slice-count warning", "T2/ADC/HIGHB depths below 3 trigger warnings, not a hard failure.", "CARPL should impose stronger completeness rules before the job is accepted."],
    ], widths=[1.3, 3.0, 2.25])
    doc.add_heading("Recommended pre-transfer integrity checks", level=2)
    add_bullets(doc, [
        "All instances share Patient ID/context, Study Instance UID, Series Instance UID, Modality=MR and Frame of Reference UID as appropriate.",
        "Unique SOP Instance UIDs and unique/expected slice positions; no duplicate pixel planes unless explicitly understood.",
        "Consistent Rows, Columns, PixelSpacing and ImageOrientationPatient within each series.",
        "A complete, monotonic, approximately uniform stack with plausible anatomical coverage.",
        "Use conventional single-frame MR Image Storage for this integration. Enhanced multi-frame MR has not been established as an equivalent input to this loader/volume path.",
        "Preserve the original pixel data, transfer syntax and spatial tags; do not recompress or rewrite geometry solely to reduce storage.",
    ])

    doc.add_heading("6. End-to-end inference pipeline", level=1)
    add_table(doc, ["Stage", "Inputs", "Processing", "Outputs / dependency"], [
        ["DICOM ingestion", "Case folder", "Recursive load; group study/series; regex select; HIGHB max-b filter; construct volumes.", "Three in-memory images: T2, ADC, HIGHB."],
        ["Organ segmentation", "T2 only", "MONAI transforms: channel-first, RAS, 1.0 mm isotropic, non-zero channel-wise normalisation. Sliding-window inference ROI 128×128×16, 50% overlap.", "Original-space prostate mask from organ/model.ts."],
        ["Organ cleanup", "Organ mask", "Keep largest 3-D connected component, fill holes, preserve class labels by nearest retained label.", "cleaned_organ.nii.gz plus cleanup metrics."],
        ["Lesion preprocessing", "T2 + ADC + HIGHB + cleaned organ", "ADC/HIGHB aligned to T2; all image channels resampled to 0.5 mm; 32-voxel-margin prostate ROI; per-channel z-score.", "3-channel prostate-centred tensor."],
        ["Lesion segmentation", "3-channel tensor", "Five RRUNet3D folds in parallel; probability maps averaged; multiply by cleaned organ; threshold 0.63.", "Fold maps, merged probability and lesion mask."],
        ["Classification", "T2 + ADC + HIGHB + organ + lesion", "0.5 mm resampling; 64³ lesion crops; 3-D ResNet; class index + 2 → PI-RADS 2–5; length rule can promote 4→5.", "lesions.txt and DICOM SR."],
        ["DICOM export", "Masks + exact selected T2 source", "Create organ/lesion/combined RTSTRUCTs referenced to T2; copy selected DICOM series.", "Viewer-importable DICOM outputs and audit artefacts."],
    ], widths=[1.15, 1.35, 3.25, 1.2], font_size=7.5)
    add_callout(doc, "Why all three channels are mandatory", "Organ inference can start from T2 alone, but lesion segmentation and lesion classification explicitly receive T2, ADC and HIGHB. Missing inputs raise operator errors or leave incomplete/empty outputs. Substituting one diffusion channel for another is not an acceptable fallback.", kind="risk")

    page_break(doc)
    doc.add_heading("7. Software, libraries, models and hardware", level=1)
    doc.add_heading("Runtime libraries pinned by the application", level=2)
    add_table(doc, ["Component", "Pinned / declared version", "Role in this workflow"], [
        ["Python", "3.10 in Ubuntu 22.04 container", "Application runtime."],
        ["PyTorch", "2.6.0, CUDA 12.4 wheel", "GPU/CPU tensors and organ, lesion and classifier inference."],
        ["MONAI", "1.4.0", "Medical-image transforms, resampling, MetaTensor and model inference utilities."],
        ["MONAI Deploy App SDK", "2.0.0", "Operator graph, DICOM loader, series selector, series-to-volume converter and application context."],
        ["Holoscan + CLI", "2.9.0", "Execution/package ecosystem used by MONAI Deploy/MAP workflows."],
        ["pydicom", "3.0.1", "DICOM header parsing, b-value extraction, UID/reference handling and output validation."],
        ["highdicom", "0.27.0", "Standards-based DICOM Comprehensive SR creation."],
        ["SimpleITK", "2.4.1", "ADC/HIGHB-to-T2 geometric resampling and NIfTI I/O."],
        ["nibabel", "5.3.2", "NIfTI affine/orientation access and file output."],
        ["scikit-image", "0.25.2", "Volume resize, connected-component properties and lesion measurements."],
        ["SciPy", ">=1.9.0", "Connected components, hole filling, distance transforms and geometry utilities."],
        ["rt-utils", "1.2.7", "RTSTRUCT construction from selected T2 and masks; locally post-processed for contour behavior."],
        ["Matplotlib", ">=3.8,<4", "Visualisation/review utilities; not the core selector."],
    ], widths=[1.25, 1.55, 3.95], font_size=7.6)
    doc.add_heading("Model artefact contract", level=2)
    add_table(doc, ["Path under /models", "Purpose", "Input contract"], [
        ["organ/model.ts", "TorchScript organ segmentation", "One T2 channel; 1.0 mm MONAI preprocessing."],
        ["fold0…fold4/model_best_foldN.pth.tar", "Five-fold RR-UNet lesion ensemble", "Three ordered channels: T2, ADC, HIGHB; 0.5 mm preprocessing."],
        ["classifier/model_best.pth.tar", "3-D ResNet PI-RADS classifier", "T2, ADC, HIGHB, organ and lesion-derived crops."],
    ], widths=[2.25, 2.05, 2.45])
    doc.add_heading("Container and hardware", level=2)
    add_bullets(doc, [
        "Base image: nvidia/cuda:12.4.1-cudnn-runtime-ubuntu22.04; PyTorch is installed from the CUDA 12.4 wheel index.",
        "app.yaml declares 8 CPU cores, one GPU, 7 GiB host memory and 7 GiB GPU memory. Separately, README.md recommends an NVIDIA GPU with at least 12 GB. For CARPL sizing, benchmark the exact production image/model set and use the higher proven requirement with headroom.",
        "The runtime uses a 1 GiB shared-memory allocation in the Docker helpers and NVIDIA Container Toolkit for GPU exposure.",
        "CPU fallback exists but is expected to be much slower; production throughput should be qualified on the actual RAPIER BOX GPU.",
        "Models, input and output are volume-mounted. Replacing source code without rebuilding/reloading the Docker image does not update a running deployment.",
    ])

    doc.add_heading("8. Storage-efficient CARPL hand-off design", level=1)
    add_text(doc, "The app can accept a full study, but that is operationally inefficient and less deterministic. The preferred architecture performs metadata-only classification at the CARPL/PACS boundary, transfers only validated payloads, and keeps immutable provenance records outside the transient RAPIER workspace.")
    add_table(doc, ["Layer", "Keep", "Retention / storage recommendation"], [
        ["PACS / source archive", "Authoritative complete study", "Remain the clinical source of truth; no change proposed."],
        ["CARPL selection record", "Study/series/SOP UIDs, descriptions, ruleset version, tag snapshot, b-value counts, rejection/ambiguity reason", "Small durable audit record; never rely only on folder names."],
        ["RAPIER input staging", "Exactly three selected series; HIGHB reduced to max-b only after validation", "Transient, per-job, read-only mount. Delete according to an approved recoverable retention policy only after output transfer and audit completion."],
        ["RAPIER output", "NIfTI intermediates, probabilities, masks, reports, RTSTRUCT/SR and selected DICOM copies", "Decide which artefacts CARPL imports versus archives. output/dicom duplicates staged source images by design."],
        ["Viewer / clinical review", "T2 evidence plus derived RTSTRUCT/SR; optionally ADC/HIGHB for context", "Use the exact selected T2 series referenced by the derived objects."],
    ], widths=[1.35, 3.15, 2.25], font_size=7.7)
    doc.add_heading("Proposed payload manifest", level=2)
    add_code(doc, "job_id, patient_context, accession_number, StudyInstanceUID\nruleset_id, app_image_digest, model_manifest_digest\nT2:    SeriesInstanceUID, SeriesDescription, SOPInstanceUID count, geometry summary\nADC:   SeriesInstanceUID, SeriesDescription, SOPInstanceUID count, geometry summary\nHIGHB: SeriesInstanceUID, SeriesDescription, source b-value histogram, retained max b-value, retained SOP count\nvalidation: zero/multiple matches, missing tags, duplicate positions, inconsistent geometry, cross-modality study/FOR checks")
    add_callout(doc, "Do not optimise away the T2 source images", "RTSTRUCT contours and the DICOM SR refer back to the selected T2 evidence. A metadata-only T2 object or lossy reconstruction cannot replace the exact source series used for inference and review.", kind="warn")

    page_break(doc)
    doc.add_heading("9. Customisation method and change control", level=1)
    add_text(doc, "Customisation should be treated as a versioned clinical-data interface change, even when the edit looks like a small regex adjustment. A wider expression may route an unintended sequence into a model without causing a software exception.")
    add_numbered(doc, [
        ("Inventory first. ", "Extract per-series metadata across the target scanners and protocol versions. The repository script fill_series_description_template.py records UIDs, SeriesDescription, ProtocolName, ImageType, geometry, slice counts and exhaustive b-value histograms."),
        ("Obtain labels. ", "Have a qualified reviewer assign exactly one role—T2, ADC, HIGHB or ignore—to each unique naming/protocol pattern, including repeats, derived maps, localizers and incomplete exams."),
        ("Edit a candidate ruleset. ", "Keep Modality=MR. Prefer explicit positive markers plus explicit exclusions. Avoid silently making ImageType mandatory for SGH unless missing-value handling is designed."),
        ("Test offline. ", "Run validate_selection_rules.py against the labelled corpus; inspect every false positive, false negative and multi-match. Do not accept aggregate accuracy alone."),
        ("Test whole studies. ", "Verify zero/one/multiple candidate counts per study and deterministic selection. The current validation CSV is row-level and does not by itself prove one-candidate-per-study behavior."),
        ("Run end to end. ", "Use representative DICOM from each scanner/vendor. Confirm selected UIDs, HIGHB retained counts, spatial alignment, input NIfTIs, masks, RTSTRUCT references and DICOM SR evidence."),
        ("Version and deploy. ", "Record the ruleset, app image digest and model manifest together. Rebuild and reload the immutable Docker image; preserve rollback."),
        ("Monitor drift. ", "Alert on unseen descriptions, missing b-value tags, multiple matches, unusual slice counts and geometry deviations."),
    ])
    doc.add_heading("Where to customise", level=2)
    add_table(doc, ["Need", "Primary location", "Caution"], [
        ["Change T2/ADC/HIGHB description matching", "Rules_T2, Rules_ADC, Rules_HIGHB in prostate_mri_lesion_seg_app/app.py", "Changes take effect only after source execution or container rebuild/reload."],
        ["Support another vendor b-value tag", "get_dicom_bvalue() in highb_filter_operator.py", "Validate VR/VM, units, per-frame semantics and mixed-tag precedence."],
        ["Reject ambiguous candidates", "Add an explicit preflight or custom selector operator", "Current default is first match; this is the highest-priority hardening opportunity."],
        ["Strengthen geometry gates", "CARPL preflight and/or a new in-app validation operator before volume conversion", "Current code warns/falls back in several cases to preserve execution."],
        ["Change model preprocessing", "organ_seg_operator.py and custom lesion/classifier operators", "Requires model compatibility and clinical performance evaluation; not a routing-only change."],
    ], widths=[1.5, 2.85, 2.4])

    doc.add_heading("10. Validation, monitoring and failure handling", level=1)
    add_table(doc, ["Condition", "Current behavior", "Production-facing response"], [
        ["No match for one role", "Downstream operator may not run; output can be empty/incomplete, and the app may still exit zero in some paths.", "Fail the job visibly; report missing T2/ADC/HIGHB and the examined descriptions."],
        ["Multiple matches for one role", "First match is selected; no ambiguity report.", "Quarantine or apply a separately validated deterministic policy; log all candidates."],
        ["HIGHB has no recognised b-value tags", "Matched series passes through unfiltered.", "Flag for manual/vendor-specific handling; never assume description b-value equals per-instance content."],
        ["Missing IOP/IPP on some slices", "Those slices are removed before stacking.", "Reject incomplete series or prove that remaining stack is complete and correct."],
        ["Fewer than 3 slices", "Warning only; affine fallback may occur.", "Reject as inadequate 3-D input unless a validated exception exists."],
        ["Empty organ mask", "Lesion inference is skipped and an all-zero lesion mask is written.", "Distinguish from a true negative; classify as model/input failure requiring review."],
        ["Container exit zero but artefacts absent", "Possible; batch wrapper inspects files to classify outcome.", "Use artefact/status checks, not process exit code alone."],
    ], widths=[1.45, 2.65, 2.65], font_size=7.7)
    doc.add_heading("Available repository guardrails", level=2)
    add_bullets(doc, [
        "scripts/validate_selection_rules.py — row-level precision/recall/F1 and error listings against labelled series descriptions.",
        "scripts/fill_series_description_template.py — recursive DICOM metadata inventory with all-slice b-value aggregation and match flags.",
        "scripts/load_and_run.py — validates Docker/model paths, runs a single case and checks required output artefacts.",
        "scripts/batch_load_and_run.py plus scripts/case_status.py — per-case isolation, audit CSV and explicit SUCCESS/FAIL vocabulary.",
        "scripts/validate_rtstruct_outputs.py — validates DICOM references, contour geometry and mask round trips.",
        "scripts/compare_batch_outputs.py — regression comparison and DICOM SR/T2 study checks.",
    ])

    page_break(doc)
    doc.add_heading("11. CARPL implementation checklist", level=1)
    add_table(doc, ["#", "Acceptance item", "Evidence CARPL should return"], [
        ["1", "One patient context and one MR Study Instance UID per job.", "Job manifest and rejection of mixed-study input."],
        ["2", "Exactly one T2, one ADC and one HIGHB candidate after the deployed expressions.", "All candidate Series UIDs/descriptions and final decision."],
        ["3", "No silent first-match behavior when multiple candidates exist.", "Ambiguity status or a documented, validated tie-break policy."],
        ["4", "Complete T2 and ADC series; complete maximum-b spatial volume for HIGHB.", "SOP counts, slice-position checks and b-value histogram."],
        ["5", "Standard (0018,9087) then Siemens (0019,100C) b-value precedence implemented exactly, including rounding.", "Per-instance extraction audit and retained/discarded counts."],
        ["6", "Missing/unknown b-value metadata is not silently guessed upstream.", "Quarantine/manual-review status and source tag snapshot."],
        ["7", "Patient, study, series, SOP and Frame of Reference identifiers preserved.", "Pre/post transfer UID comparison."],
        ["8", "Rows/Columns, PixelSpacing, orientation, positions and transfer integrity preserved.", "Geometry consistency report and checksums where available."],
        ["9", "RAPIER input mounted read-only and output isolated per job.", "Mount configuration and job-specific paths."],
        ["10", "Success based on required artefacts and audit status, not container exit code alone.", "Final status plus artefact inventory."],
        ["11", "Derived RTSTRUCT/SR returned with the exact selected T2 evidence/reference.", "DICOM reference validation report."],
        ["12", "Ruleset, app container and model set versioned together with rollback.", "Ruleset ID, image digest and model deployment manifest/checksums."],
    ], widths=[0.35, 3.9, 2.5], font_size=7.6)
    add_callout(doc, "Suggested go-live gate", "Demonstrate a representative acceptance set spanning scanner manufacturers, protocol variants, repeats, multi-b acquisitions, missing-tag cases and deliberately ambiguous studies. CARPL and SGH should jointly sign off the selected UIDs, retained HIGHB instances, spatial alignment and returned DICOM references.", kind="safe")

    doc.add_heading("Appendix A. Exact deployed regular expressions", level=1)
    doc.add_heading("T2", level=2)
    add_code(doc, r"^(?!.*(?:obl|oblq|oblique|obq|cor|sag|pelvis|whole[_ ]?pelvis|wholepelvis|tirm|stir|space|localizer|survey|t1|\+c|adc|dwi|diff|tracew?|_fs_|_fs\b|\bfs_|\bfs\b))(?=.*t2)(?=.*(?:tra|\bax\b|axial|^ax[_ ]|_ax[_ ]|AX[_ ])).*")
    doc.add_heading("ADC", level=2)
    add_code(doc, r"^(?!.*(?:localizer|survey))(?=.*adc).*")
    doc.add_heading("HIGHB", level=2)
    add_code(doc, r"^(?!.*adc)(?!.*(?:localizer|survey))(?:(?=.*tracew?)|(?=.*dwi)|(?=.*diff)|(?=.*calc)|(?=.*compute[dr]?)|(?=.*(?:focus|muse).*b[0-9])).*")
    add_callout(doc, "Matching semantics", "MONAI first attempts case-insensitive exact string equality for string attributes and then case-insensitive regex search. These anchored expressions are therefore effectively evaluated case-insensitively over the whole SeriesDescription.", kind="info")

    doc.add_heading("Appendix B. Required tags and preservation rules", level=1)
    add_table(doc, ["Tag", "Keyword / purpose", "Use in current code or integration"], [
        ["(0008,0060)", "Modality", "Must be MR for all three selectors."],
        ["(0008,103E)", "SeriesDescription", "Primary T2/ADC/HIGHB routing field."],
        ["(0020,000D)", "StudyInstanceUID", "Study grouping and T2 evidence relationship; CARPL should enforce one study."],
        ["(0020,000E)", "SeriesInstanceUID", "Series grouping, selected-series copying and exact T2 lookup for RTSTRUCT."],
        ["(0008,0018)", "SOPInstanceUID", "Instance identity and derived-object evidence references."],
        ["(0020,0052)", "FrameOfReferenceUID", "Must be preserved and checked for spatial compatibility; not an app selector condition."],
        ["(0018,9087)", "DiffusionBValue", "Preferred per-instance HIGHB pruning tag."],
        ["(0019,100C)", "Siemens private b-value", "Fallback when standard b-value tag is unavailable."],
        ["(0020,0037)", "ImageOrientationPatient", "Slice normal, distance and volume affine; missing slice is removed."],
        ["(0020,0032)", "ImagePositionPatient", "Slice distance and ordering; missing slice is removed."],
        ["(0028,0030)", "PixelSpacing", "In-plane geometry; ImagerPixelSpacing is a loader fallback."],
        ["(0020,0013)", "InstanceNumber", "Fallback ordering only; not a substitute for spatial position."],
        ["(0020,0010)", "StudyID", "May be absent in SGH de-id; RTSTRUCT path creates a temporary copy populated from AccessionNumber."],
        ["(0008,0050)", "AccessionNumber", "Used only for temporary StudyID repair when required; preserve genuine values/blanks."],
    ], widths=[1.05, 2.05, 3.65], font_size=7.6)

    page_break(doc)
    doc.add_heading("Appendix C. Source-code traceability map", level=1)
    add_table(doc, ["Topic", "Authoritative repository source"], [
        ["Pipeline graph, deployed selectors and SGH volume-converter patch", "prostate_mri_lesion_seg_app/app.py"],
        ["Per-instance highest-b pruning", "prostate_mri_lesion_seg_app/highb_filter_operator.py"],
        ["T2 preprocessing and organ inference", "prostate_mri_lesion_seg_app/organ_seg_operator.py"],
        ["T2/ADC/HIGHB save, alignment, ROI, five-fold lesion inference, threshold and selected DICOM copies", "prostate_mri_lesion_seg_app/custom_lesion_seg_operator.py"],
        ["Lesion classification and reporting", "prostate_mri_lesion_seg_app/custom_lesion_classifier_operator.py"],
        ["Organ cleanup and shared array utilities", "prostate_mri_lesion_seg_app/common.py"],
        ["RTSTRUCT construction and T2 references", "prostate_mri_lesion_seg_app/rtstruct_utils.py"],
        ["DICOM SR construction", "prostate_mri_lesion_seg_app/dicom_sr_utils.py"],
        ["Python dependency pins", "prostate_mri_lesion_seg_app/requirements.txt and docker/requirements.txt"],
        ["Container runtime", "docker/Dockerfile and docker/docker-compose.yaml"],
        ["Resource declaration", "prostate_mri_lesion_seg_app/app.yaml"],
        ["Rule validation", "scripts/validate_selection_rules.py and series_description_sgh_deid_data.tnt_with_anh_input.csv"],
        ["Series metadata / b-value inventory", "scripts/fill_series_description_template.py"],
        ["Deployment and per-case status", "scripts/load_and_run.py, scripts/batch_load_and_run.py and scripts/case_status.py"],
        ["High-level operational documentation", "README.md and DOCKER GUIDE README.md"],
    ], widths=[2.75, 4.0], font_size=7.8)
    doc.add_heading("Document assumptions and factual boundaries", level=2)
    add_bullets(doc, [
        "app.py and the current checked-out operator files are treated as the deployed behavior described here. Historical app backup files are not treated as active.",
        "The retrospective selection metrics were independently re-run while preparing this document; they test the annotated CSV’s fields and rule semantics, not pixel quality or study-level uniqueness.",
        "Recommendations labelled CARPL or production-facing are integration hardening proposals. They are intentionally distinguished from current in-box behavior.",
        "No inference code, filters, models or deployment configuration were changed to create this document.",
    ])
    add_callout(doc, "Owner fields for CARPL/SGH to complete", "Ruleset owner: __________  |  CARPL implementation version: __________  |  RAPIER image digest: __________  |  Model manifest digest: __________  |  Acceptance set/date: __________  |  Approvers: __________", kind="info")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
